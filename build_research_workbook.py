#!/usr/bin/env python3
"""生成「天气-股市」长期研究 Excel 工作簿。"""

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "天气与股市关系研究记录表.xlsx"
CSV_OUTPUT = ROOT / "每日数据记录.csv"
DATA_SHEET = "每日数据"
CHART_SHEET = "走势图"
STAT_SHEET = "统计检验"
GUIDE_SHEET = "使用说明"

INITIAL_ROW = {
    "date": date(2026, 5, 20),
    "weather": {"北京": 7, "上海": 4, "深圳": 5},
    "market": {"北京": 3, "上海": 4, "深圳": 5},  # 北证50 / 上证指数 / 深证成指
    "note": "研究首日：天气与股市情绪不完全同步的样本日",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAX_DATA_ROW = 2000
HEADER_ROW = 1
FIRST_DATA_ROW = 2


def style_header(ws, row: int, cols: int, fill=None):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill or HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def build_daily_sheet(wb: Workbook):
    """第 1 张表：常规「一行 = 一天」数据总表。"""
    ws = wb.active
    ws.title = DATA_SHEET
    ws.sheet_view.showGridLines = True

    # 第 1 行即为表头（打开即可见标准表格）
    headers = [
        "日期",
        "北京天气(1-10)",
        "上海天气(1-10)",
        "深圳天气(1-10)",
        "北京股市(北证50)",
        "上海股市(上证指数)",
        "深圳股市(深证成指)",
        "天气均值",
        "股市均值",
        "备注",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=HEADER_ROW, column=col, value=h)
    style_header(ws, HEADER_ROW, len(headers))

    widths = [13, 11, 11, 11, 11, 11, 11, 10, 10, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[HEADER_ROW].height = 36

    # 第 2 行：首日数据
    r = FIRST_DATA_ROW
    ws.cell(row=r, column=1, value=INITIAL_ROW["date"]).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=2, value=INITIAL_ROW["weather"]["北京"])
    ws.cell(row=r, column=3, value=INITIAL_ROW["weather"]["上海"])
    ws.cell(row=r, column=4, value=INITIAL_ROW["weather"]["深圳"])
    ws.cell(row=r, column=5, value=INITIAL_ROW["market"]["北京"])
    ws.cell(row=r, column=6, value=INITIAL_ROW["market"]["上海"])
    ws.cell(row=r, column=7, value=INITIAL_ROW["market"]["深圳"])
    ws.cell(row=r, column=10, value=INITIAL_ROW["note"])

    # 均值列：仅当该行有日期时计算
    for row in range(FIRST_DATA_ROW, MAX_DATA_ROW + 1):
        ws.cell(
            row=row,
            column=8,
            value=f'=IF($A{row}="","",AVERAGE(B{row}:D{row}))',
        )
        ws.cell(
            row=row,
            column=9,
            value=f'=IF($A{row}="","",AVERAGE(E{row}:G{row}))',
        )
        if row > FIRST_DATA_ROW:
            ws.cell(row=row, column=1).number_format = "yyyy-mm-dd"

    dv = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="10",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="评分范围错误",
        error="请填写 1 到 10 之间的整数。",
    )
    dv.sqref = f"B{FIRST_DATA_ROW}:G{MAX_DATA_ROW}"
    ws.add_data_validation(dv)

    table_ref = f"A{HEADER_ROW}:J{MAX_DATA_ROW}"
    tab = Table(displayName="DailyScores", ref=table_ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=True,
    )
    ws.add_table(tab)

    ws.conditional_formatting.add(
        f"B{FIRST_DATA_ROW}:G{MAX_DATA_ROW}",
        ColorScaleRule(
            start_type="num",
            start_value=1,
            start_color="F8696B",
            mid_type="num",
            mid_value=5.5,
            mid_color="FFEB84",
            end_type="num",
            end_value=10,
            end_color="63BE7B",
        ),
    )

    ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 115

    # 顶部提示（放在表右侧，不挡住数据表）
    ws["L1"] = "← 左侧即「每日数据总表」：每行一天"
    ws["L1"].font = Font(bold=True, color="C00000", size=11)
    ws["L2"] = "在表格最下方空行续填即可"
    ws["L3"] = "折线图见「走势图」标签页"

    return ws


def build_chart_sheet(wb: Workbook):
    ws_data = wb[DATA_SHEET]
    ws = wb.create_sheet(CHART_SHEET)

    def line_chart(title, y_title, cols, names):
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = y_title
        chart.x_axis.title = "日期"
        chart.style = 10
        chart.height = 14
        chart.width = 24
        cats = Reference(
            ws_data,
            min_col=1,
            min_row=FIRST_DATA_ROW,
            max_row=MAX_DATA_ROW,
        )
        for col, name in zip(cols, names):
            data = Reference(
                ws_data,
                min_col=col,
                min_row=HEADER_ROW,
                max_row=MAX_DATA_ROW,
            )
            chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        for i, name in enumerate(names):
            if i < len(chart.series):
                chart.series[i].title = SeriesLabel(v=name)
        return chart

    ws.add_chart(
        line_chart(
            "三城天气评分（按日）",
            "天气 1-10",
            [2, 3, 4],
            ["北京天气", "上海天气", "深圳天气"],
        ),
        "A1",
    )
    ws.add_chart(
        line_chart(
            "三城股市评分（按日）",
            "股市 1-10",
            [5, 6, 7],
            ["北京股市(北证50)", "上海股市(上证)", "深圳股市(深证)"],
        ),
        "A18",
    )
    ws.add_chart(
        line_chart(
            "天气均值 vs 股市均值",
            "评分 1-10",
            [8, 9],
            ["天气均值", "股市均值"],
        ),
        "A35",
    )
    ws["A50"] = "说明：上图数据来自「每日数据」表，保存工作簿后自动更新。"
    return ws


def build_stat_sheet(wb: Workbook):
    ws = wb.create_sheet(STAT_SHEET)
    ws["A1"] = "假设检验（H₀：同城天气与同城股市无线性相关；α=0.05）"
    ws["A1"].font = Font(bold=True, size=12, color="1F4E79")
    ws.merge_cells("A1:H1")

    ws["A3"] = "有效样本量 n（已填写日期的行数）"
    ws["B3"] = f"=COUNTA('{DATA_SHEET}'!$A${FIRST_DATA_ROW}:$A${MAX_DATA_ROW})"
    ws["A4"] = "说明"
    ws["B4"] = "建议 n≥30 再解读；n<3 显示「数据不足」。同城配对：北京↔北证、上海↔上证、深圳↔深证。"

    pairs = [
        ("北京：天气 ↔ 股市", "B", "E"),
        ("上海：天气 ↔ 股市", "C", "F"),
        ("深圳：天气 ↔ 股市", "D", "G"),
        ("天气三城均值 ↔ 股市三城均值", "H", "I"),
    ]

    header_row = 6
    cols = ["检验配对", "Pearson r", "t 统计量", "双侧 p 值", "α=0.05 结论", "|r| 效应量"]
    for c, h in enumerate(cols, 1):
        ws.cell(row=header_row, column=c, value=h)
    style_header(ws, header_row, len(cols))

    r0 = header_row + 1
    for i, (label, wx, mx) in enumerate(pairs):
        row = r0 + i
        ws.cell(row=row, column=1, value=label)
        wcol = f"'{DATA_SHEET}'!${wx}${FIRST_DATA_ROW}:${wx}${MAX_DATA_ROW}"
        mcol = f"'{DATA_SHEET}'!${mx}${FIRST_DATA_ROW}:${mx}${MAX_DATA_ROW}"
        r_cell = f"B{row}"
        ws.cell(row=row, column=2, value=f'=IF($B$3<3,"—",CORREL({wcol},{mcol}))')
        ws.cell(
            row=row,
            column=3,
            value=f'=IF($B$3<3,"—",ABS({r_cell})*SQRT($B$3-2)/SQRT(1-{r_cell}^2))',
        )
        ws.cell(row=row, column=4, value=f'=IF($B$3<3,"—",T.DIST.2T(C{row},$B$3-2))')
        ws.cell(
            row=row,
            column=5,
            value=(
                f'=IF($B$3<3,"数据不足",IF(D{row}<0.05,'
                f'"拒绝 H₀（p<0.05）","不能拒绝 H₀"))'
            ),
        )
        ws.cell(
            row=row,
            column=6,
            value=(
                f'=IF($B$3<3,"—",IF(ABS({r_cell})<0.3,"弱",'
                f'IF(ABS({r_cell})<0.5,"中等","较强")))'
            ),
        )

    matrix_row = r0 + len(pairs) + 2
    ws.cell(row=matrix_row, column=1, value="3×3 相关矩阵（行=天气，列=股市）").font = Font(bold=True)
    w_letters = ["B", "C", "D"]
    m_letters = ["E", "F", "G"]
    cities = ["北京", "上海", "深圳"]
    ws.cell(row=matrix_row + 1, column=1, value="")
    for j, c in enumerate(cities, 2):
        ws.cell(row=matrix_row + 1, column=j, value=f"{c}股市")
    for i, (city, wl) in enumerate(zip(cities, w_letters)):
        mr = matrix_row + 2 + i
        ws.cell(row=mr, column=1, value=f"{city}天气")
        for j, ml in enumerate(m_letters, 2):
            wref = f"'{DATA_SHEET}'!${wl}${FIRST_DATA_ROW}:${wl}${MAX_DATA_ROW}"
            mref = f"'{DATA_SHEET}'!${ml}${FIRST_DATA_ROW}:${ml}${MAX_DATA_ROW}"
            ws.cell(row=mr, column=j, value=f'=IF($B$3<3,"—",CORREL({wref},{mref}))')

    ws.column_dimensions["A"].width = 28
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20
    return ws


def build_guide_sheet(wb: Workbook):
    ws = wb.create_sheet(GUIDE_SHEET)
    lines = [
        "使用说明",
        "",
        "【数据表在哪？】",
        "  → 打开文件后，第一个标签页就是「每日数据」。",
        "  → 第 1 行是表头，从第 2 行起：一行 = 一个交易日。",
        "  → 列顺序：日期 | 北京/上海/深圳 天气 | 北京/上海/深圳 股市 | 均值 | 备注",
        "",
        "【股市列对应哪个指数？】",
        "  北京股市 = 北证50 | 上海股市 = 上证指数 | 深圳股市 = 深证成指",
        "",
        "【每日怎么填？】",
        "  在表格最下方空行填写 1–10 整数，保存即可。",
        "",
        "【图表 / 统计】",
        "  「走势图」= 折线图 | 「统计检验」= 相关与 p 值",
        "",
        "【纯文本备份】",
        f"  同文件夹下的 {CSV_OUTPUT.name} 可用记事本直接查看。",
    ]
    for i, text in enumerate(lines, 1):
        cell = ws.cell(row=i, column=1, value=text)
        if i == 1:
            cell.font = Font(bold=True, size=14, color="1F4E79")
    ws.column_dimensions["A"].width = 70


def export_csv():
    headers = [
        "日期",
        "北京天气",
        "上海天气",
        "深圳天气",
        "北京股市",
        "上海股市",
        "深圳股市",
        "备注",
    ]
    with CSV_OUTPUT.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerow(
            [
                INITIAL_ROW["date"].isoformat(),
                INITIAL_ROW["weather"]["北京"],
                INITIAL_ROW["weather"]["上海"],
                INITIAL_ROW["weather"]["深圳"],
                INITIAL_ROW["market"]["北京"],
                INITIAL_ROW["market"]["上海"],
                INITIAL_ROW["market"]["深圳"],
                INITIAL_ROW["note"],
            ]
        )


def main():
    wb = Workbook()
    build_daily_sheet(wb)
    build_chart_sheet(wb)
    build_stat_sheet(wb)
    build_guide_sheet(wb)
    wb.save(OUTPUT)
    export_csv()
    print(f"已生成: {OUTPUT}")
    print(f"已生成: {CSV_OUTPUT}")


if __name__ == "__main__":
    main()
